"""Run using uv with python-docx and pdfplumber; all fixtures live in temporary dirs."""
from __future__ import annotations
import copy
import json
import os
import re
from pathlib import Path
import shutil
import struct
import subprocess
import sys
import unicodedata
import tempfile
import time
import unittest
from unittest.mock import patch
import zlib

sys.dont_write_bytecode=True
ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT/'scripts'))
import md_export as exporter
import presets
import ooxml as ox
import checks
from docx import Document
from docx.oxml.ns import qn


def png(path,width=900,height=220):
    def chunk(kind,data): return struct.pack('>I',len(data))+kind+data+struct.pack('>I',zlib.crc32(kind+data)&0xffffffff)
    data=b''.join(b'\x00'+b'\xdd\xdd\xdd'*width for _ in range(height))
    path.write_bytes(b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',width,height,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(data))+chunk(b'IEND',b''))


def mutate(path,change):
    package=ox.read_package(path); parts=ox.xml_parts(package)
    change(parts); ox.write_package(path,package,parts)


def run_script(*args):
    return subprocess.run([sys.executable,str(ROOT/'scripts/md_export.py'),*map(str,args)],capture_output=True,text=True,timeout=240,env={**os.environ,'PYTHONDONTWRITEBYTECODE':'1'})


class PresetTests(unittest.TestCase):
    def test_sizes(self):
        for name,n in presets.SIZES.items(): self.assertEqual(presets.size_pt(name),n)
        self.assertEqual(presets.size_pt('10.5pt'),10.5)
        for value in ['oops','nan','inf','0','-1','10.3']:
            with self.assertRaises(presets.ConfigError): presets.size_pt(value)

    def test_complete_json(self):
        for name in ('general','academic','gongwen'):
            p=presets.load_preset(name)
            self.assertEqual(len(p['styles']),34)
            for s in p['styles'].values(): self.assertEqual(set(s),presets.FIELDS)
        self.assertFalse(presets.load_preset('gongwen')['styles']['Heading 3']['bold'])

    def test_overrides(self):
        p=presets.load_preset('academic',['body.size=四号','body.line_spacing=fixed:22','heading2.bold=false'],orientation='landscape')
        self.assertEqual(p['styles']['Body Text']['size'],14)
        self.assertEqual(p['styles']['First Paragraph']['line_spacing'],'fixed:22')
        self.assertEqual(p['styles']['Verbatim Char']['size'],14)
        self.assertFalse(p['styles']['Heading 2']['bold'])
        self.assertEqual(p['page']['orientation'],'landscape')

    def test_invalid_overrides_list_keys(self):
        for value in ['body.nope=2','body.bold=yes','body.size=wat','body.line_spacing=22','title.color=FF0000','page.margins_cm=[50,50,50,50]']:
            with self.assertRaisesRegex(presets.ConfigError,'可用键'): presets.load_preset(overrides=[value])

    def test_custom_json(self):
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'custom.json'; p=presets.load_preset(); p['styles']['Title']['size']=20
            path.write_text(json.dumps(p)); self.assertEqual(presets.load_preset(str(path))['styles']['Title']['size'],20)
            del p['styles']['Normal']['color']; path.write_text(json.dumps(p))
            with self.assertRaises(presets.ConfigError): presets.load_preset(str(path))

    def test_invalid_json_structure(self):
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'bad.json'
            for value in (None,[],{'name':'bad'},dict(presets.load_preset(),styles=None)):
                path.write_text(json.dumps(value))
                with self.assertRaises(presets.ConfigError): presets.load_preset(str(path))

    def test_font_fallback_and_degradation(self):
        p=presets.load_preset('gongwen')
        resolved,decisions=presets.resolve_fonts(p,{'simsun','simhei','fangsong','kaiti','times new roman','courier new'})
        self.assertEqual(resolved['styles']['Title']['eastAsia'],'SimSun')
        self.assertTrue(resolved['styles']['Title']['bold'])
        self.assertEqual(resolved['styles']['Heading 3']['eastAsia'],'FangSong')
        self.assertFalse(resolved['styles']['Heading 3']['bold'])
        self.assertTrue(next(d for d in decisions if d['required']=='FZXiaoBiaoSong-B05S')['degraded'])

    def test_strict_and_unknown_fonts(self):
        for available in (None,set()):
            with self.assertRaises(presets.ConfigError): presets.resolve_fonts(presets.load_preset(),available,True)
        p,ds=presets.resolve_fonts(presets.load_preset(),None)
        self.assertTrue(all(d['status']=='unknown' for d in ds))
        p,ds=presets.resolve_fonts(presets.load_preset(),set())
        self.assertTrue(all(d['status']=='missing' for d in ds))
        self.assertEqual(p['styles']['Body Text']['eastAsia'],'SimSun')

    def test_linux_fallback(self):
        p,ds=presets.resolve_fonts(presets.load_preset('gongwen'),{'Noto Serif CJK SC','Noto Sans CJK SC','DejaVu Serif','DejaVu Sans Mono'})
        self.assertTrue(all(d['status'] in ('fallback','pass') for d in ds))
        self.assertEqual(p['styles']['Body Text']['eastAsia'],'Noto Serif CJK SC')

    def test_invisible_characters_ignored_in_text_match(self):
        # Linux 版 LibreOffice 会丢掉 Pandoc 写进表格的零宽空格
        self.assertEqual(checks.norm('a\u200b b\u00ad\ufeffc'),'abc')
        self.assertFalse(checks.visible('\u200b')); self.assertTrue(checks.visible('中'))

    def test_fontconfig_wrapper(self):
        with tempfile.TemporaryDirectory() as t:
            config=Path(t)/'fonts.conf'; config.touch(); wrapper=Path(t)/'soffice'
            wrapper.write_text('#!/bin/sh\n: "${FONTCONFIG_FILE:='+str(config)+'}"\n')
            with patch.dict(os.environ,{},clear=True): self.assertEqual(presets.font_environment(wrapper)['FONTCONFIG_FILE'],str(config))

    def test_macos_homebrew_fontconfig(self):
        with tempfile.TemporaryDirectory() as t:
            config=Path(t)/'fonts.conf'; config.touch(); missing=Path(t)/'missing.conf'
            with patch.dict(os.environ,{},clear=True), patch.object(presets,'MAC_FONTCONFIG',(missing,config)):
                with patch.object(presets.sys,'platform','darwin'): self.assertEqual(presets.font_environment()['FONTCONFIG_FILE'],str(config))
                with patch.object(presets.sys,'platform','linux'): self.assertNotIn('FONTCONFIG_FILE',presets.font_environment())
            # 用户自己设的优先
            with patch.dict(os.environ,{'FONTCONFIG_FILE':'/custom/fonts.conf'},clear=True), patch.object(presets,'MAC_FONTCONFIG',(config,)), patch.object(presets.sys,'platform','darwin'):
                self.assertEqual(presets.font_environment()['FONTCONFIG_FILE'],'/custom/fonts.conf')


class AstAndCliTests(unittest.TestCase):
    def header(self,text,level=1,classes=None): return {'t':'Header','c':[level,['',classes or [],[]],[{'t':'Str','c':text}]]}
    def ast(self,title,body):
        return dict(meta={'title':{'t':'MetaInlines','c':[{'t':'Str','c':title}]}} if title else {},blocks=body)

    def test_title_same(self):
        ast,flags,removed=exporter.prepare_ast(self.ast('标题',[self.header('标题'),self.header('正文')]),presets.load_preset())
        self.assertEqual(removed,'标题'); self.assertEqual(len(ast['blocks']),1)
        self.assertIn('title',ast['meta'])

    def test_title_different_and_only_one(self):
        for title,body in [('标题',[self.header('不同')]),('标题',[]),('',[self.header('正文')])]:
            ast,flags,removed=exporter.prepare_ast(self.ast(title,body),presets.load_preset())
            self.assertIsNone(removed); self.assertEqual(len(ast['blocks']),len(body))

    def test_handwritten_patterns(self):
        for kind,level,text in [('academic',1,'1 概述'),('academic',2,'1.1 工作'),('academic',3,'1.1.1 方法'),('gongwen',1,'一、概述'),('gongwen',1,'一、 概述'),('gongwen',2,'（一）工作'),('gongwen',3,'1.方法'),('gongwen',4,'（1）细节')]: self.assertTrue(exporter.handwritten(text,level,kind))
        for kind in ('academic','gongwen'):
            self.assertFalse(exporter.handwritten('2026 年计划',1,kind))
            self.assertFalse(exporter.handwritten('1.1版本',2,kind))
        self.assertFalse(exporter.handwritten('1.1 方法',3,'gongwen'))

    def test_per_heading_disable(self):
        p=presets.load_preset('gongwen')
        ast,flags,removed=exporter.prepare_ast(self.ast('',[self.header('一、概述'),self.header('工作'),self.header('特殊',2,['unnumbered'])]),p)
        self.assertEqual([f['disabled'] for f in flags],[True,False,True])

    def test_output_inference(self):
        md=Path('/tmp/test.md')
        for argv,expected in [([],('docx',['/tmp/test.docx'])),(['--to','pdf'],('pdf',['/tmp/test.pdf'])),(['--output','/tmp/new.pdf'],('pdf',['/tmp/new.pdf'])),(['--output','/tmp/new.PDF'],('pdf',['/tmp/new.PDF'])),(['--output','/tmp/new.docx','--to','both'],('both',['/tmp/new.docx','/tmp/new.pdf']))]:
            args=exporter.parse_args([str(md),*argv]); to,paths=exporter.targets(args,md)
            self.assertEqual((to,[str(p) for p in paths]),expected)
        with self.assertRaises(exporter.ConversionError): exporter.targets(exporter.parse_args([str(md),'--to','docx','--output','/tmp/new.pdf']),md)
        with self.assertRaises(exporter.ConversionError): exporter.targets(exporter.parse_args(['/tmp/same.docx','--output','/tmp/same.docx']),Path('/tmp/same.docx'))

    def test_soffice_precedence(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'s'; p.write_text('#!/bin/sh\n'); p.chmod(0o755)
            with patch.dict(os.environ,{'MD_EXPORT_SOFFICE':str(p)}):
                self.assertEqual(exporter.find_soffice(),str(p))
                with self.assertRaises(exporter.ConversionError): exporter.find_soffice('/does-not-exist')
            self.assertIsNone(exporter.find_soffice('/does-not-exist',required=False))


class XmlTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(); cls.t=Path(cls.temp.name)
        cls.default=cls.t/'default.docx'
        if shutil.which('pandoc'): exporter.run(['pandoc','-o',cls.default,'--print-default-data-file','reference.docx'],30)
        else: Document().save(cls.default)
    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()
    def reference(self,name='general'):
        path=self.t/(name+'.docx'); p=presets.load_preset(name); ox.build_reference(self.default,path,p)
        return path,p,ox.xml_parts(ox.read_package(path))

    def test_reference_fonts_themes_black_nonitalic(self):
        for name in ('general','academic','gongwen'):
            path,p,parts=self.reference(name)
            self.assertFalse(checks.color_scan(parts))
            styles=parts['word/styles.xml']
            for s in styles.findall('w:style',ox.NS):
                self.assertFalse(any(checks.truth(i) for i in s.iter(qn('w:i'))))
            for r in styles.iter(qn('w:rFonts')):
                # All explicitly written styles and defaults have four families.
                for k in ox.THEME_FONTS: self.assertNotIn(qn('w:'+k),r.attrib)
            for n in p['styles']:
                sid=Document(path).styles[n].style_id
                r=styles.find(f'w:style[@w:styleId="{sid}"]/w:rPr/w:rFonts',ox.NS)
                self.assertEqual(set(r.attrib),{qn('w:'+a) for a in ('ascii','hAnsi','cs','eastAsia')})

    def test_reference_body_indent_spacing_and_a4(self):
        path,p,parts=self.reference('academic'); r=checks.Resolver(parts)
        for name in ('Body Text','First Paragraph'):
            v,s=r.style(Document(path).styles[name].style_id)
            self.assertEqual(v['first_line_chars'],2); self.assertEqual(v['line_spacing'],'fixed:20'); self.assertEqual(v['size'],12)
        self.assertEqual(r.style('Normal')[0]['first_line_chars'],0)
        size=parts['word/document.xml'].find('.//w:pgSz',ox.NS)
        self.assertAlmostEqual(int(size.get(qn('w:w')))/20,595.28,delta=1)
        self.assertAlmostEqual(int(size.get(qn('w:h')))/20,841.89,delta=1)

    def test_reference_footer_and_grid(self):
        path,p,parts=self.reference('gongwen')
        self.assertIsNotNone(parts['word/settings.xml'].find('w:evenAndOddHeaders',ox.NS))
        self.assertEqual(parts['word/document.xml'].find('.//w:docGrid',ox.NS).get(qn('w:linePitch')),'560')
        footers=[root for name,root in parts.items() if name.startswith('word/footer')]
        self.assertEqual(len(footers),2)
        self.assertEqual({f.find('.//w:jc',ox.NS).get(qn('w:val')) for f in footers},{'right','left'})
        for f in footers: self.assertEqual(f.find('.//w:fldSimple',ox.NS).get(qn('w:instr')),'PAGE')

    def test_numbering_definition(self):
        for name,texts in [('academic',['%1','%1.%2','%1.%2.%3']),('gongwen',['%1、','（%2）','%3.','（%4）'])]:
            path,p,parts=self.reference(name)
            headings=[n for n in parts['word/numbering.xml'].findall('w:abstractNum',ox.NS) if n.find('w:lvl/w:pStyle',ox.NS) is not None]
            self.assertEqual([n.get(qn('w:val')) for n in headings[-1].findall('w:lvl/w:lvlText',ox.NS)],texts)
            self.assertEqual(checks.Resolver(parts).style('Heading1')[0]['bold'],name=='academic')

    def make_table(self,rows=3,heads=1,merged=False):
        doc=Document(); tab=doc.add_table(rows=rows,cols=2)
        for i,row in enumerate(tab.rows):
            for j,cell in enumerate(row.cells): cell.text=f'row{i}cell{j}'; ox.child(ox.child(cell._tc,'tcPr'),'tcBorders').append(ox.el('left',val='double',sz=24))
            if i<heads: ox.child(row._tr,'trPr').append(ox.el('tblHeader',val=1))
        if merged: tab.cell(heads-1,0).merge(tab.cell(heads-1,1))
        return tab._tbl

    def test_grid_borders_clear_cell_overrides(self):
        t=self.make_table(); ox.table_borders(t,'grid')
        self.assertEqual(len(t.find('w:tblPr/w:tblBorders',ox.NS)),6)
        self.assertTrue(all(n.get(qn('w:sz'))=='4' for n in t.find('w:tblPr/w:tblBorders',ox.NS)))
        self.assertFalse(t.findall('w:tr/w:tc/w:tcPr/w:tcBorders',ox.NS))

    def test_three_line_last_header_and_merged(self):
        t=self.make_table(4,2); ox.table_borders(t,'three_line')
        self.assertEqual(len(t.find('w:tblPr/w:tblBorders',ox.NS)),2)
        self.assertFalse(t.findall('w:tr',ox.NS)[0].findall('w:tc/w:tcPr/w:tcBorders',ox.NS))
        self.assertEqual(len(t.findall('w:tr',ox.NS)[1].findall('w:tc/w:tcPr/w:tcBorders/w:bottom',ox.NS)),2)
        self.assertEqual(len(t.findall('w:tr/w:trPr/w:tblHeader',ox.NS)),2)

    def test_three_line_single_and_no_header(self):
        for rows,heads in [(1,1),(1,0),(3,0)]:
            t=self.make_table(rows,heads); ox.table_borders(t,'three_line')
            self.assertFalse(t.findall('w:tr/w:tc/w:tcPr/w:tcBorders',ox.NS))
            self.assertEqual(len(t.find('w:tblPr/w:tblBorders',ox.NS)),2)

    def test_three_line_all_header_and_merged_cells(self):
        for rows,heads,merged in [(2,2,False),(4,2,True)]:
            t=self.make_table(rows,heads,merged); ox.table_borders(t,'three_line')
            last=t.findall('w:tr',ox.NS)[heads-1]
            cells=last.findall('w:tc',ox.NS)
            self.assertEqual(len(cells),1 if merged else 2)
            self.assertEqual(len(last.findall('w:tc/w:tcPr/w:tcBorders/w:bottom',ox.NS)),len(cells))
            self.assertTrue(all(tc.find('w:tcPr/w:tcBorders/w:bottom',ox.NS).get(qn('w:sz'))=='6' for tc in cells))
            if merged: self.assertEqual(cells[0].find('w:tcPr/w:gridSpan',ox.NS).get(qn('w:val')),'2')

    def test_effective_chain_toggle_and_character_style(self):
        path,p,parts=self.reference(); resolver=checks.Resolver(parts)
        root=parts['word/styles.xml']; parent=ox.el('style',type='paragraph',styleId='Parent')
        parent.append(ox.el('name',val='Parent')); ox.child(parent,'rPr').append(ox.el('b',val=1)); root.append(parent)
        derived=ox.el('style',type='paragraph',styleId='Derived'); derived.append(ox.el('name',val='Derived')); derived.append(ox.el('basedOn',val='Parent')); ox.child(derived,'rPr').append(ox.el('b',val=1)); root.append(derived)
        resolver=checks.Resolver(parts); self.assertFalse(resolver.style('Derived')[0]['bold'])
        para=ox.el('p'); ox.child(para,'pPr').append(ox.el('pStyle',val='BodyText')); run=ox.el('r'); ox.child(run,'rPr').append(ox.el('rStyle',val='VerbatimChar')); para.append(run)
        v,s=resolver.effective(para,run); self.assertEqual(v['ascii'],'Courier New'); self.assertEqual(s['ascii'],'character style:VerbatimChar')
        ox.child(run,'rPr').append(ox.el('sz',val=38)); self.assertEqual(resolver.effective(para,run)[0]['size'],19)
        ox.child(para,'pPr').append(ox.el('jc',val='right')); self.assertEqual(resolver.effective(para,run)[0]['alignment'],'right')

    def test_effective_table_condition_and_numbering(self):
        path,p,parts=self.reference('academic')
        table_style=parts['word/styles.xml'].find('w:style[@w:styleId="Table"]',ox.NS)
        cond=ox.el('tblStylePr',type='firstRow'); ox.child(cond,'rPr').append(ox.el('sz',val=40)); table_style.append(cond)
        table=self.make_table(); ox.child(table,'tblPr').append(ox.el('tblStyle',val='Table'))
        para=table.find('w:tr/w:tc/w:p',ox.NS)
        # Remove paragraph/default size so the conditional table layer is visible.
        normal=parts['word/styles.xml'].find('w:style[@w:styleId="Normal"]/w:rPr/w:sz',ox.NS); normal.getparent().remove(normal)
        r=checks.Resolver(parts); self.assertEqual(r.effective(para)[0]['size'],20)
        h=ox.el('p'); ox.child(h,'pPr').append(ox.el('pStyle',val='Heading1'))
        v,s=r.effective(h); self.assertEqual(v['size'],15); self.assertEqual(s['size'],'numbering level')

    def test_color_scan_numbering_header_footnote(self):
        path,p,parts=self.reference('academic')
        for part in ('word/numbering.xml','word/footer1.xml'):
            ox.child(parts[part],'rPr').append(ox.el('color',val='FF0000',themeColor='accent1'))
        self.assertEqual(len(checks.color_scan(parts)),2)

    def test_unknown_style(self):
        path,p,parts=self.reference(); para=ox.el('p'); ox.child(para,'pPr').append(ox.el('pStyle',val='Undefined'))
        with self.assertRaises(ValueError): checks.Resolver(parts).effective(para)


class TransactionTests(unittest.TestCase):
    def test_both_rollback_existing_outputs(self):
        with tempfile.TemporaryDirectory() as t:
            t=Path(t); a=t/'a'; b=t/'b'; out1=t/'out1'; out2=t/'out2'
            a.write_bytes(b'new1'); b.write_bytes(b'new2'); out1.write_bytes(b'old1'); out2.write_bytes(b'old2')
            real_replace=os.replace; count=0
            def fail_second(src,dest):
                nonlocal count
                count+=1
                if count==2: raise OSError('simulated second publication failure')
                return real_replace(src,dest)
            with patch.object(exporter.os,'replace',side_effect=fail_second):
                with self.assertRaises(OSError): exporter.publish([(a,out1),(b,out2)])
            self.assertEqual(out1.read_bytes(),b'old1'); self.assertEqual(out2.read_bytes(),b'old2')
            self.assertEqual({x.name for x in t.iterdir()},{'a','b','out1','out2'})

    def test_publication_and_rollback_failure_preserves_backup(self):
        with tempfile.TemporaryDirectory() as t:
            t=Path(t); a=t/'a'; b=t/'b'; one=t/'one'; two=t/'two'
            for path,text in ((a,'new1'),(b,'new2'),(one,'old1'),(two,'old2')): path.write_text(text)
            real=os.replace; count=0
            def fail(src,dest):
                nonlocal count
                count+=1
                if count in (2,3): raise OSError('publication failure' if count==2 else 'rollback failure')
                return real(src,dest)
            with patch.object(exporter.os,'replace',side_effect=fail):
                with self.assertRaises(OSError) as caught: exporter.publish([(a,one),(b,two)])
            backups=list(t.glob('*.bak'))
            self.assertEqual(len(backups),1,'failed rollback must retain the old output backup')
            self.assertEqual(backups[0].read_text(),'old1')
            self.assertIn('publication failure',str(caught.exception))
            recovery=caught.exception.recovery
            self.assertEqual(recovery[0]['status'],'fail')
            self.assertEqual(Path(recovery[0]['backup']),backups[0])
            self.assertIn('rollback failure',recovery[0]['error'])
            self.assertEqual(two.read_text(),'old2')
            self.assertFalse(list(t.glob('*.tmp')))

    def test_both_rollback_new_first(self):
        with tempfile.TemporaryDirectory() as t:
            t=Path(t); a=t/'a'; b=t/'b'; a.touch(); b.touch(); real=os.replace
            def replace(src,dest):
                if Path(dest).name=='two': raise OSError('fail')
                return real(src,dest)
            with patch.object(exporter.os,'replace',side_effect=replace):
                with self.assertRaises(OSError): exporter.publish([(a,t/'one'),(b,t/'two')])
            self.assertFalse((t/'one').exists())

    def test_rollback_continues_after_one_restore_fails(self):
        with tempfile.TemporaryDirectory() as t:
            t=Path(t); files=[]
            for i in range(3):
                src=t/f'new{i}'; dest=t/f'out{i}'
                src.write_text(f'new{i}'); dest.write_text(f'old{i}'); files.append((src,dest))
            real=os.replace; count=0
            def fail(src,dest):
                nonlocal count
                count+=1
                if count in (3,4): raise OSError('injected I/O failure')
                return real(src,dest)
            with patch.object(exporter.os,'replace',side_effect=fail):
                with self.assertRaises(OSError) as caught: exporter.publish(files)
            self.assertEqual([r['status'] for r in caught.exception.recovery],['fail','pass'])
            self.assertEqual(files[0][1].read_text(),'old0')
            self.assertEqual(files[2][1].read_text(),'old2')
            backups=list(t.glob('*.bak')); self.assertEqual(len(backups),1)
            self.assertEqual(backups[0].read_text(),'old1')

    def test_ambiguous_same_text_note_is_unknown_not_wrong_size(self):
        p=presets.load_preset(); width,height=[v*72/25.4 for v in presets.PAPER_MM['A4']]
        paragraphs=[]
        for part,style in [('word/document.xml','First Paragraph'),('word/footnotes.xml','Footnote Text')]:
            paragraphs.append(dict(location=part+'/p1/'+style,text='Same',style=style,part=part,has_math=False,note_marker='1',runs=[dict(text='Same',required=p['styles'][style])]))
        glyphs=[]
        for top,size in [(80,12),(700,9)]:
            for i,letter in enumerate('Same'):
                glyphs.append(dict(text=letter,x0=100+i*size,x1=100+(i+1)*size,top=top,bottom=top+size,size=size,fontname='TimesNewRomanPSMT',non_stroking_color=0))
        from types import SimpleNamespace
        fake=SimpleNamespace(pages=[SimpleNamespace(width=width,height=height,chars=glyphs)])
        with patch.object(checks.pdfplumber,'open') as opened:
            opened.return_value.__enter__.return_value=fake
            result=checks.check_pdf('unused',p,dict(paragraphs=paragraphs),[])
        formats=[c for c in result['checks'] if c['location'].endswith('/格式')]
        self.assertEqual([c['status'] for c in formats],['unknown','unknown'])
        self.assertTrue(all(c['pdf'] is None for c in formats))

    def test_body_superscript_is_not_a_footnote_location(self):
        p=presets.load_preset(); width,height=[v*72/25.4 for v in presets.PAPER_MM['A4']]
        paragraphs=[dict(location=part+'/p1/'+style,text='Same',style=style,part=part,has_math=False,note_marker='1',runs=[dict(text='Same',required=p['styles'][style])]) for part,style in [('word/document.xml','First Paragraph'),('word/footnotes.xml','Footnote Text')]]
        glyphs=[dict(text='1',x0=90,x1=94,top=81,bottom=86,size=5,fontname='TimesNewRomanPSMT',non_stroking_color=0)]
        for top,size in [(80,12),(700,9)]:
            for i,letter in enumerate('Same'):
                glyphs.append(dict(text=letter,x0=100+i*size,x1=100+(i+1)*size,top=top,bottom=top+size,size=size,fontname='TimesNewRomanPSMT',non_stroking_color=0))
        from types import SimpleNamespace
        fake=SimpleNamespace(pages=[SimpleNamespace(width=width,height=height,chars=glyphs)])
        with patch.object(checks.pdfplumber,'open') as opened:
            opened.return_value.__enter__.return_value=fake
            result=checks.check_pdf('unused',p,dict(paragraphs=paragraphs),[])
        formats=[c for c in result['checks'] if c['location'].endswith('/格式')]
        self.assertEqual([c['status'] for c in formats],['unknown','unknown'])

    def test_timeout_kills_process_tree(self):
        if os.name!='posix': self.skipTest('POSIX process group assertion')
        with tempfile.TemporaryDirectory() as t:
            t=Path(t); fake=t/'soffice'; marker=t/'child-survived'
            fake.write_text('#!/bin/sh\n(sleep 0.6; touch "'+str(marker)+'") &\nwait\n'); fake.chmod(0o755)
            start=time.monotonic()
            with self.assertRaisesRegex(exporter.ConversionError,'超时'): exporter.run([fake],0.1)
            self.assertLess(time.monotonic()-start,0.6); time.sleep(0.65)
            self.assertFalse(marker.exists())

    def test_empty_pdf_output(self):
        with tempfile.TemporaryDirectory() as t:
            t=Path(t); fake=t/'soffice'; fake.write_text('#!/bin/sh\nexit 0\n'); fake.chmod(0o755); source=t/'doc.docx'; source.touch()
            with self.assertRaisesRegex(exporter.ConversionError,'未生成 PDF'): exporter.libreoffice(source,t,fake,10,os.environ.copy())

    def test_pdf_black_and_font_matching(self):
        for c in [None,0,(0,0,0),(0,0,0,1)]: self.assertTrue(checks.black(c))
        for c in [1,(1,0,0),(0,0,0,0)]: self.assertFalse(checks.black(c))
        self.assertTrue(checks.font_matches('BAAAAA+TimesNewRomanPSMT','Times New Roman'))
        self.assertTrue(checks.font_matches('CAAAAA+STSongtiSC-Regular','Songti SC'))
        self.assertFalse(checks.font_matches('LiberationSerif','SimSun'))


class IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try: soffice=exporter.find_soffice()
        except exporter.ConversionError: soffice=None
        if not shutil.which('pandoc') or not soffice:
            if os.environ.get('MD_EXPORT_REQUIRE_TOOLS')=='1': raise AssertionError('MD_EXPORT_REQUIRE_TOOLS=1: pandoc 或 soffice 缺失')
            raise unittest.SkipTest('pandoc 或 soffice 缺失')
        cls.soffice=soffice
        cls.tmp=tempfile.TemporaryDirectory(prefix='md-export-tests-'); cls.t=Path(cls.tmp.name)
        cls.artifacts=Path(os.environ['MD_EXPORT_TEST_ARTIFACTS']) if os.environ.get('MD_EXPORT_TEST_ARTIFACTS') else cls.t/'artifacts'
        cls.artifacts.mkdir(parents=True,exist_ok=True)
        cls.generated={}
    @classmethod
    def tearDownClass(cls): cls.tmp.cleanup()

    def sample(self,name):
        t=self.artifacts/name; t.mkdir(exist_ok=True); (t/'img').mkdir(exist_ok=True); png(t/'img/relative.png')
        md=t/'sample.md'
        table='| 序号 | 名称 | 数值 | 描述 |\n|---|---|---|---|\n'+''.join(f'| {i:02} | 中文字段 {i:02} | {i*3} | 表格跨页验收 {i:02} English |\n' for i in range(48))
        code='\n'.join(f'print("长代码跨页 line {i:02} 中文", "'+('repeat_'*14)+'")' for i in range(34))
        md.write_text('''---
title: 综合样例 Export Verification
---

# 综合样例 Export Verification

中英文正文 Mixed text。含 **加粗**、`inline code` 和 [链接 Link](https://example.com)。

# 一级标题 Heading One

2026 年计划作为正文。

## 二级标题 Heading Two

### 三级标题 Heading Three

#### 四级标题 Heading Four

##### 五级标题 Heading Five

列表示例：

- 第一条 item A
- 第二条 item B

> 引用文字 Quote。

| 单行表头 | Header |
|---|---|

: 单行表题

![相对路径大图](img/relative.png){width=24cm}

行内公式 $\\frac{a}{b}$ 结束。

$$\\frac{x^2}{y}=z$$

脚注引用[^a]。

[^a]: 中文脚注 footnote。

重复文本 Repeat text。

重复文本 Repeat text。

'''+table+'\n: 跨页宽表 Table Caption\n\n```python\n'+code+'\n```\n\n末尾正文 Final paragraph。\n',encoding='utf-8')
        return md

    def convert_sample(self,name):
        md=self.sample(name); proc=run_script(md,'--preset',name,'--to','both')
        self.assertEqual(proc.returncode,0,proc.stderr+'\n'+proc.stdout[:500]+self.failure_summary(proc.stdout))
        report=json.loads(proc.stdout); (md.parent/'report.json').write_text(proc.stdout)
        self.assertEqual(report['conversion'],'pass'); self.assertEqual(report['docx_check'],'pass'); self.assertEqual(report['pdf_check'],'pass')
        self.assertEqual(report['visual'],'not_done'); self.assertTrue(report['published'])
        self.assertGreaterEqual(report['pages'],3)
        parts=ox.xml_parts(ox.read_package(md.with_suffix('.docx')))
        self.assertTrue(any(x.startswith('word/media/') for x in ox.read_package(md.with_suffix('.docx'))))
        self.assertIsNotNone(parts['word/document.xml'].find('.//m:f',ox.NS))
        self.assertIsNotNone(parts['word/document.xml'].find('.//m:oMathPara',ox.NS))
        self.assertIsNotNone(parts['word/document.xml'].find('.//w:drawing',ox.NS))
        for table in parts['word/document.xml'].iter(qn('w:tbl')):
            for para in table.iter(qn('w:p')): self.assertEqual(ox.pstyle(para),'TableText')
        import pdfplumber
        with pdfplumber.open(md.with_suffix('.pdf')) as pdf:
            text='\n'.join(p.extract_text() or '' for p in pdf.pages)
            self.assertEqual(text.count('综合样例'),1); self.assertEqual(text.count('重复文本'),2)
            self.assertNotIn('\\frac',text); self.assertNotIn('$',text)
            self.assertGreaterEqual(sum(len(p.images) for p in pdf.pages),1)
            self.assertIn('相对路径大图',text)
            if name=='gongwen':
                self.assertIn('一、',text); self.assertIn('（一）',text)
                self.assertIsNotNone(parts['word/settings.xml'].find('w:evenAndOddHeaders',ox.NS))
                self.assertIsNotNone(parts['word/document.xml'].find('.//w:footerReference[@w:type="even"]',ox.NS))
        code_checks=[c for c in report['checks'] if c['location'].endswith('/Source Code/格式')]
        self.assertTrue(code_checks)
        for c in code_checks:
            for run in c['pdf']['runs']:
                for glyph in run['glyphs']: self.assertAlmostEqual(glyph['size'],10,delta=0.65)
        for para in parts['word/document.xml'].iter(qn('w:p')):
            if ox.pstyle(para)=='SourceCode': self.assertFalse(para.findall('.//w:rPr/w:rStyle',ox.NS))
        self.generated[name]=(md,report)
        return md,report

    @staticmethod
    def failure_summary(stdout):
        try:
            r=json.loads(stdout)
            return '\n'+str([(c['location'],c.get('details'),c.get('pdf',{})) for c in r.get('checks',[]) if c['status'] in ('fail','unknown')])[:10000]
        except ValueError: return ''

    def test_general_integration(self): self.convert_sample('general')
    def test_academic_integration(self): self.convert_sample('academic')
    def test_gongwen_integration(self): self.convert_sample('gongwen')

    def small(self,name,content='正文中文 Body text。\n',*extra):
        t=self.t/name; t.mkdir(exist_ok=True); md=t/'input.md'; md.write_text(content,encoding='utf-8')
        proc=run_script(md,*extra); return md,proc

    def test_failed_rollback_paths_are_saved_in_diagnostic_report(self):
        t=(self.t/'failed_restore').resolve(); t.mkdir(); md=t/'input.md'; md.write_text('中文正文。')
        docx=md.with_suffix('.docx'); pdf=md.with_suffix('.pdf')
        docx.write_bytes(b'old docx'); pdf.write_bytes(b'old pdf')
        real=os.replace
        def fail(src,dest):
            if Path(dest)==pdf or Path(dest)==docx and Path(src).suffix=='.bak': raise OSError('injected publication or restore failure')
            return real(src,dest)
        with patch.object(exporter.os,'replace',side_effect=fail):
            with self.assertRaises(exporter.ConversionError) as caught:
                exporter.convert(exporter.parse_args([str(md),'--to','both','--diagnostics',str(t)]))
        report=caught.exception.report
        saved=json.loads((Path(report['diagnostics'])/'report.json').read_text())
        self.assertEqual(saved['publication'],report['publication'])
        recovery=saved['publication']['recovery']; self.assertEqual(len(recovery),1)
        self.assertEqual(recovery[0]['status'],'fail'); self.assertEqual(recovery[0]['output'],str(docx))
        self.assertEqual(Path(recovery[0]['backup']).read_bytes(),b'old docx')
        self.assertEqual(pdf.read_bytes(),b'old pdf'); self.assertFalse(report['published'])

    def test_math_in_table_keeps_table_auto_spacing(self):
        md,proc=self.small('table_math','| 表内公式 |\n|---|\n| $\\frac{X}{Y}$ |\n','--preset','academic','--to','both')
        self.assertEqual(proc.returncode,0,self.failure_summary(proc.stdout))
        parts=ox.xml_parts(ox.read_package(md.with_suffix('.docx')))
        para=next(p for p in parts['word/document.xml'].iter(qn('w:p')) if p.find('.//m:f',ox.NS) is not None)
        self.assertEqual(ox.pstyle(para),'TableText')
        self.assertEqual(checks.Resolver(parts).effective(para)[0]['line_spacing'],'multiple:1')

    def test_footer_allowed_fallback_still_reports_replacement(self):
        md,proc=self.small('footer_replacement','正文 English。\n','--to','both')
        self.assertEqual(proc.returncode,0,proc.stdout)
        report=json.loads(proc.stdout)
        p,_=presets.resolve_fonts(presets.load_preset(),{d['actual'] for d in report['font_decisions']})
        # 要求字体本身已回退到 Liberation Serif 时，造不出「改用回退字体」的场景
        if p['styles']['Page Number']['ascii']!='Times New Roman': self.skipTest('本机没有 Times New Roman')
        dx=checks.check_docx(md.with_suffix('.docx'),p)
        tampered=md.parent/'tampered.docx'; shutil.copy2(md.with_suffix('.docx'),tampered)
        def change(parts):
            style=parts['word/styles.xml'].find('w:style[@w:styleId="PageNumber"]',ox.NS)
            runs=[r for name,tree in parts.items() if name.startswith('word/footer') for r in tree.iter(qn('w:r'))]
            for element in [style]+runs:
                ox.replace(ox.child(element,'rPr'),'rFonts',ascii='Liberation Serif',hAnsi='Liberation Serif',cs='Liberation Serif',eastAsia=p['styles']['Page Number']['eastAsia'])
        mutate(tampered,change)
        pdf,_=exporter.libreoffice(tampered,md.parent,self.soffice,120,presets.font_environment(self.soffice))
        ordinary=checks.check_pdf(pdf,p,dx,report['font_decisions'])
        self.assertEqual(ordinary['status'],'pass')
        self.assertTrue(any('页码字体替换' in w for w in ordinary['warnings']))
        strict=checks.check_pdf(pdf,p,dx,report['font_decisions'],True)
        self.assertEqual(strict['status'],'fail')
        self.assertTrue(any(c['status']=='fail' and '页码' in c['location'] for c in strict['checks']))

    def test_inline_images_preserve_body_and_heading_styles(self):
        t=self.t/'inline_images'; t.mkdir(); png(t/'pic.png',20,10)
        md=t/'input.md'
        md.write_text('# 标题 ![](pic.png){width=0.2cm}\n\n正文有小图 ![](pic.png){width=0.2cm}，后面有文字。\n\n独立图之前。\n\n![](pic.png){width=1cm}\n')
        proc=run_script(md,'--preset','academic','--to','both')
        self.assertEqual(proc.returncode,0,proc.stdout+self.failure_summary(proc.stdout))
        report=json.loads(proc.stdout)
        p,_=presets.resolve_fonts(presets.load_preset('academic'),{d['actual'] for d in report['font_decisions']})
        parts=ox.xml_parts(ox.read_package(md.with_suffix('.docx')))
        paras=list(parts['word/document.xml'].iter(qn('w:p')))
        heading=next(p for p in paras if ox.ptext(p).startswith('标题'))
        body=next(p for p in paras if ox.ptext(p).startswith('正文有小图'))
        self.assertEqual(ox.pstyle(heading),'Heading1')
        self.assertIn(ox.pstyle(body),('BodyText','FirstParagraph'))
        resolver=checks.Resolver(parts)
        values,_=resolver.effective(body)
        self.assertEqual(values['line_spacing'],'fixed:20'); self.assertEqual(values['first_line_chars'],2)
        self.assertIsNotNone(resolver.number_level(heading))
        self.assertTrue(any(ox.pstyle(p)=='Figure' for p in paras if not ox.ptext(p) and p.find('.//w:drawing',ox.NS) is not None))

    def test_docx_heading_count_and_levels_match_ast(self):
        md,proc=self.small('heading_structure','# 一级\n\n## 二级\n\n正文。','--preset','academic')
        self.assertEqual(proc.returncode,0,proc.stdout)
        report=json.loads(proc.stdout)
        p,_=presets.resolve_fonts(presets.load_preset('academic'),{d['actual'] for d in report['font_decisions']})
        flags=[dict(level=1,disabled=False),dict(level=2,disabled=False)]
        original=md.with_suffix('.docx')
        for name,style in [('missing','Figure'),('level','Heading3')]:
            dest=md.parent/(name+'.docx'); shutil.copy2(original,dest)
            def change(parts):
                heading=next(p for p in parts['word/document.xml'].iter(qn('w:p')) if ox.pstyle(p)=='Heading2')
                ox.replace(ox.child(heading,'pPr'),'pStyle',val=style)
            mutate(dest,change)
            result=checks.check_docx(dest,p,flags)
            self.assertTrue(any(c['location']=='标题层级' and c['status']=='fail' for c in result['checks']),[(c['location'],c['status']) for c in result['checks'] if c['status']!='pass'])

    def test_body_and_footnote_identical_text_real_pdf(self):
        md,proc=self.small('note_overlap','重复的内容。\n\n第二个段落带脚注[^a]。\n\n[^a]: 重复的内容。\n','--to','both')
        self.assertEqual(proc.returncode,0,self.failure_summary(proc.stdout))
        report=json.loads(proc.stdout)
        self.assertEqual(report['pdf_check'],'pass')
        formats=[c for c in report['checks'] if c['location'].endswith('/格式') and c.get('pdf') and any('重复的内容' in r['text'] for r in c['pdf'].get('runs',[]))]
        self.assertEqual(len(formats),2)
        for c in formats:
            sizes={g['size'] for r in c['pdf']['runs'] for g in r['glyphs']}
            self.assertEqual(sizes,{9.0} if 'footnotes' in c['location'] else {12.0})

    def test_inline_code_bold_cli_and_custom_preset(self):
        for name,opts in [('cli',['--set','inline_code.bold=true']),('custom',[])]:
            if name=='custom':
                custom=self.t/'bold.json'; p=presets.load_preset(); p['styles']['Verbatim Char']['bold']=True
                custom.write_text(json.dumps(p)); opts=['--preset',str(custom)]
            md,proc=self.small('inline_bold_'+name,'这是 `code` 示例。\n\n# 标题 `code`\n',*opts)
            self.assertEqual(proc.returncode,0,self.failure_summary(proc.stdout))
            report=json.loads(proc.stdout)
            self.assertTrue(report['published'])
            code_runs=[r for c in report['checks'] if c['location'].startswith('word/document.xml/') for r in (c.get('docx') or {}).get('runs',[]) if r['text']=='code']
            self.assertEqual([r['effective']['bold'] for r in code_runs],[True,False])
            self.assertEqual([r['required']['bold'] for r in code_runs],[True,False])

    def test_academic_inline_fraction_minimum_spacing_real_pdf(self):
        content='前一行 PREV。\n分式 $\\frac{X}{Y}$ 结束。\n后一行 NEXT。\n\n普通段落。\n'
        md,proc=self.small('inline_fraction',content,'--preset','academic','--to','both')
        self.assertEqual(proc.returncode,0,self.failure_summary(proc.stdout))
        report=json.loads(proc.stdout)
        parts=ox.xml_parts(ox.read_package(md.with_suffix('.docx')))
        para=next(p for p in parts['word/document.xml'].iter(qn('w:p')) if p.find('.//m:f',ox.NS) is not None)
        space=para.find('w:pPr/w:spacing',ox.NS)
        self.assertIsNotNone(space,'inline OMML needs a paragraph spacing override')
        self.assertEqual(space.get(qn('w:lineRule')),'atLeast'); self.assertEqual(space.get(qn('w:line')),'400')
        self.assertIn(ox.pstyle(para),('BodyText','FirstParagraph'))
        exceptions=report['line_spacing_exceptions']
        self.assertEqual(len(exceptions),1); self.assertIn('分式',exceptions[0]['text'])
        self.assertEqual(exceptions[0]['line_spacing'],'at_least:20')
        ordinary=next(p for p in report['checks'] if p['location'].startswith('word/document.xml/') and any('普通段落' in r['text'] for r in (p.get('docx') or {}).get('runs',[])))
        self.assertEqual(ordinary['docx']['effective']['line_spacing'],'fixed:20')
        tampered=md.parent/'exact.docx'; shutil.copy2(md.with_suffix('.docx'),tampered)
        def remove_exception(parts):
            para=next(p for p in parts['word/document.xml'].iter(qn('w:p')) if p.find('.//m:f',ox.NS) is not None)
            para.find('w:pPr/w:spacing',ox.NS).set(qn('w:lineRule'),'exact')
        mutate(tampered,remove_exception)
        p,_=presets.resolve_fonts(presets.load_preset('academic'),{d['actual'] for d in report['font_decisions']})
        self.assertEqual(checks.check_docx(tampered,p)['status'],'fail')
        import pdfplumber
        with pdfplumber.open(md.with_suffix('.pdf')) as pdf:
            chars=pdf.pages[0].chars
            stream=''.join(c['text'] for c in chars)
            # 公式字母在有的平台是数学斜体码位（如 U+1D44B），按 NFKC 归一后再找
            letter=lambda want: [c for c in chars if unicodedata.normalize('NFKC',c['text'])==want]
            # Linux 版 LibreOffice 把公式画成图形，PDF 里没有公式字符，只能核对上面的 DOCX 行距
            if not (letter('X') or letter('Y')): self.skipTest(f'PDF 里的公式不是文字，无法按字符位置测裁切：{stream!r}')
            self.assertTrue(letter('X') and letter('Y'),f'PDF 中找不到分式字母：{stream!r}')
            numerator=letter('X')[0]; denominator=letter('Y')[0]
            previous=chars[stream.index('PREV'):stream.index('PREV')+4]
            following=chars[stream.index('NEXT'):stream.index('NEXT')+4]
            self.assertGreaterEqual(numerator['top'],max(c['bottom'] for c in previous)-0.5)
            self.assertGreaterEqual(denominator['top'],numerator['bottom']-0.5)
            self.assertLessEqual(denominator['bottom'],min(c['top'] for c in following)+0.5)

    def test_docx_only_without_soffice(self):
        md,proc=self.small('only','中文正文。\n','--soffice','/does-not-exist')
        self.assertEqual(proc.returncode,0,proc.stdout); r=json.loads(proc.stdout)
        self.assertEqual(r['docx_check'],'pass'); self.assertEqual(r['pdf_check'],'unknown'); self.assertIsNone(r['soffice'])
        self.assertTrue(md.with_suffix('.docx').exists()); self.assertFalse(md.with_suffix('.pdf').exists())

    def test_title_variants_real(self):
        for n,content,title_count,h1_count in [('different','---\ntitle: 文档标题\n---\n\n# 正文标题\n\n正文。',1,1),('only_yaml','---\ntitle: 文档标题\n---\n\n正文。',1,0),('only_h1','# 正文标题\n\n正文。',0,1)]:
            md,proc=self.small(n,content)
            self.assertEqual(proc.returncode,0,proc.stdout)
            root=ox.xml_parts(ox.read_package(md.with_suffix('.docx')))['word/document.xml']
            self.assertEqual(sum(ox.pstyle(p)=='Title' for p in root.iter(qn('w:p'))),title_count)
            self.assertEqual(sum(ox.pstyle(p)=='Heading1' for p in root.iter(qn('w:p'))),h1_count)

    def test_numbered_and_unnumbered_mix_real(self):
        md,proc=self.small('manual','# 一、手写标题\n\n# 自动标题\n\n## 特殊 {.unnumbered}\n\n## 自动二级\n','--preset','gongwen','--to','both')
        self.assertEqual(proc.returncode,0,proc.stdout+self.failure_summary(proc.stdout))
        r=json.loads(proc.stdout); self.assertEqual(len(r['numbering_disabled']),2)
        ps=[p for p in ox.xml_parts(ox.read_package(md.with_suffix('.docx')))['word/document.xml'].iter(qn('w:p')) if ox.pstyle(p).startswith('Heading')]
        self.assertEqual([p.find('w:pPr/w:numPr/w:numId',ox.NS) is not None for p in ps],[True,False,True,False])

    def test_failed_conversion_preserves_old_output_and_diagnostics(self):
        t=self.t/'failed'; t.mkdir(); fake=t/'soffice'; fake.write_text('#!/bin/sh\nexit 0\n'); fake.chmod(0o755)
        md=t/'source.md'; md.write_text('中文正文。'); old=t/'source.pdf'; old.write_bytes(b'old pdf')
        proc=run_script(md,'--to','pdf','--soffice',fake,'--diagnostics',t/'diagnostics')
        self.assertNotEqual(proc.returncode,0); r=json.loads(proc.stdout); self.assertFalse(r['published'])
        self.assertEqual(old.read_bytes(),b'old pdf'); self.assertTrue((Path(r['diagnostics'])/'document.docx').is_file())
        self.assertTrue((Path(r['diagnostics'])/'report.json').is_file())

    def test_timeout_conversion_preserves_old_output(self):
        t=self.t/'timeout_convert'; t.mkdir(); fake=t/'soffice'
        fake.write_text('#!/bin/sh\nsleep 30\n'); fake.chmod(0o755)
        md=t/'source.md'; md.write_text('中文正文。'); old=t/'source.pdf'; old.write_bytes(b'old pdf')
        proc=run_script(md,'--to','pdf','--soffice',fake,'--timeout','1','--diagnostics',t/'diagnostics')
        self.assertNotEqual(proc.returncode,0); report=json.loads(proc.stdout)
        self.assertIn('超时',report['error']); self.assertFalse(report['published'])
        self.assertEqual(old.read_bytes(),b'old pdf')
        self.assertTrue((Path(report['diagnostics'])/'document.docx').exists())

    def test_strict_fonts_failure(self):
        with patch.object(exporter,'installed_fonts',return_value=set()):
            md=self.t/'strict.md'; md.write_text('中文正文。')
            with self.assertRaisesRegex(exporter.ConversionError,'缺少字体'): exporter.convert(exporter.parse_args([str(md),'--strict-fonts']))
        self.assertFalse(md.with_suffix('.docx').exists())

    def test_invalid_cli_returns_2(self):
        for i,opts in enumerate([['--set','body.nope=1'],['--set','body.size=bad'],['--to','pdf','--output',str(self.t/'bad.docx')],['--timeout','0']]):
            md,proc=self.small('invalid'+str(i),'正文。',*opts); self.assertEqual(proc.returncode,2,proc.stdout)
            self.assertFalse(md.with_suffix('.docx').exists())

    def test_tampered_docx_effective_values_fail(self):
        md,proc=self.small('tamper','# 一级标题\n\n正文 `inline code` 文字。\n','--preset','academic')
        self.assertEqual(proc.returncode,0,proc.stdout)
        original=md.with_suffix('.docx'); report=json.loads(proc.stdout)
        p,_=presets.resolve_fonts(presets.load_preset('academic'),{d['actual'] for d in report['font_decisions']})
        self.assertEqual(checks.check_docx(original,p)['status'],'pass')
        scenarios={
            'size':lambda ps: ox.replace(ox.child(next(r for r in ps['word/document.xml'].iter(qn('w:r')) if r.find('w:t',ox.NS) is not None),'rPr'),'sz',val=60),
            'color':lambda ps: ox.replace(ox.child(next(ps['word/document.xml'].iter(qn('w:r'))),'rPr'),'color',val='FF0000'),
            'charstyle':lambda ps: ox.replace(ox.child(ps['word/styles.xml'].find('w:style[@w:styleId="VerbatimChar"]',ox.NS),'rPr'),'rFonts',ascii='Arial',hAnsi='Arial',cs='Arial',eastAsia='Arial'),
            'numbering':lambda ps: ps['word/styles.xml'].find('w:style[@w:styleId="Heading1"]/w:pPr',ox.NS).remove(ps['word/styles.xml'].find('w:style[@w:styleId="Heading1"]/w:pPr/w:numPr',ox.NS))}
        for name,change in scenarios.items():
            dest=md.parent/(name+'.docx'); shutil.copy2(original,dest); mutate(dest,change)
            result=checks.check_docx(dest,p)
            self.assertEqual(result['status'],'fail',name)

    def test_pdf_deleted_paragraph_and_repetition_fail(self):
        md,proc=self.small('delete','唯一段落 Alpha。\n\n重复段落 Repeat。\n\n重复段落 Repeat。\n','--to','both')
        self.assertEqual(proc.returncode,0,proc.stdout+self.failure_summary(proc.stdout))
        report=json.loads(proc.stdout)
        p,_=presets.resolve_fonts(presets.load_preset(),{d['actual'] for d in report['font_decisions']})
        dx=checks.check_docx(md.with_suffix('.docx'),p)
        self.assertEqual(dx['status'],'pass')
        corrupted=md.parent/'deleted.docx'; shutil.copy2(md.with_suffix('.docx'),corrupted)
        def delete(parts):
            body=parts['word/document.xml'].find('w:body',ox.NS)
            ps=body.findall('w:p',ox.NS); body.remove(ps[0]); body.remove(ps[-1])
        mutate(corrupted,delete)
        pdf,_=exporter.libreoffice(corrupted,md.parent,self.soffice,120,presets.font_environment(self.soffice))
        result=checks.check_pdf(pdf,p,dx,report['font_decisions'])
        self.assertEqual(result['status'],'fail')
        missing=[c for c in result['checks'] if c['status']=='fail' and '/内容' in c['location']]
        self.assertEqual(len(missing),2)

    def test_pdf_cjk_replaced_with_western_fails(self):
        md,proc=self.small('western','这是中文正文。\n','--to','both')
        self.assertEqual(proc.returncode,0,proc.stdout)
        report=json.loads(proc.stdout)
        p,_=presets.resolve_fonts(presets.load_preset(),{d['actual'] for d in report['font_decisions']})
        dx=checks.check_docx(md.with_suffix('.docx'),p)
        self.assertEqual(dx['status'],'pass')
        data=md.with_suffix('.pdf').read_bytes()
        import pdfplumber
        with pdfplumber.open(md.with_suffix('.pdf')) as pdf:
            font=next(c['fontname'] for pg in pdf.pages for c in pg.chars if any(checks.is_cjk(x) for x in c['text'])).encode('ascii')
        pattern=rb'/(?:BaseFont|FontName)\s*/'+re.escape(font)+rb'\b'
        corrupted,count=re.subn(pattern,lambda m:m.group().replace(font,b'AAAAAA+Times'.ljust(len(font),b' ')),data)
        self.assertGreater(count,0)
        path=md.parent/'western-font.pdf'; path.write_bytes(corrupted)
        result=checks.check_pdf(path,p,dx,report['font_decisions'])
        self.assertEqual(result['status'],'fail')
        self.assertTrue(any(c['location']=='PDF/中文缺字风险' and c['status']=='fail' for c in result['checks']))

    def test_real_pdf_wrong_size_and_color_fail(self):
        md,proc=self.small('pdf_tamper','中文正文 Western text。\n','--to','both')
        self.assertEqual(proc.returncode,0,proc.stdout)
        report=json.loads(proc.stdout)
        p,_=presets.resolve_fonts(presets.load_preset(),{d['actual'] for d in report['font_decisions']})
        original=md.with_suffix('.docx'); dx=checks.check_docx(original,p)
        self.assertEqual(dx['status'],'pass')
        for tag,attributes in [('sz',dict(val=44)),('color',dict(val='FF0000'))]:
            tmp=md.parent/tag; tmp.mkdir(); target=tmp/'corrupt.docx'; shutil.copy2(original,target)
            def change(parts):
                run=next(r for r in parts['word/document.xml'].iter(qn('w:r')) if r.find('w:t',ox.NS) is not None)
                ox.replace(ox.child(run,'rPr'),tag,**attributes)
            mutate(target,change)
            pdf,_=exporter.libreoffice(target,tmp,self.soffice,120,presets.font_environment(self.soffice))
            result=checks.check_pdf(pdf,p,dx,report['font_decisions'])
            self.assertEqual(result['status'],'fail',tag)
            failures=[c for c in result['checks'] if c['status']=='fail']
            self.assertTrue(failures)
            self.assertTrue(any('字号' in str(c) if tag=='sz' else '非黑' in str(c) for c in failures))

    def test_unknown_raw_xml_is_not_published(self):
        md,proc=self.small('raw','```{=openxml}\n<w:p><w:r><w:t>raw text</w:t></w:r></w:p>\n```\n')
        self.assertNotEqual(proc.returncode,0)
        report=json.loads(proc.stdout)
        self.assertIn(report['docx_check'],('unknown','fail'))
        self.assertFalse(report['published']); self.assertFalse(md.with_suffix('.docx').exists())

    def test_landscape_relative_image_formula(self):
        tmp=self.t/'landscape'; tmp.mkdir(); (tmp/'img').mkdir(); png(tmp/'img/p.png',200,80)
        md=tmp/'input.md'; md.write_text('中文段落。\n1）第一条\n2）第二条\n\n![图片](img/p.png)\n\n行内公式 $\\frac{a}{b}$ 结束。\n\n```tex\n\\frac{a}{b}\n```\n')
        proc=run_script(md,'--orientation','landscape','--to','both')
        self.assertEqual(proc.returncode,0,proc.stdout+self.failure_summary(proc.stdout))
        import pdfplumber
        with pdfplumber.open(md.with_suffix('.pdf')) as pdf:
            self.assertAlmostEqual(pdf.pages[0].width,841.89,delta=2)
            self.assertAlmostEqual(pdf.pages[0].height,595.28,delta=2)
            text='\n'.join(p.extract_text() or '' for p in pdf.pages)
            self.assertEqual(text.count('\\frac{a}{b}'),1)
            self.assertNotIn('$',text); self.assertIn('1）第一条',text)
            self.assertGreaterEqual(sum(len(p.images) for p in pdf.pages),1)
        doc=Document(md.with_suffix('.docx'))
        from docx.enum.section import WD_ORIENT
        self.assertEqual(doc.sections[0].orientation,WD_ORIENT.LANDSCAPE)

if __name__=='__main__': unittest.main()
